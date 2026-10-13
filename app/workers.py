"""
app.workers — Background workers (DocWorker, ThumbnailWorker)
"""
import sys, os, io, re, json, csv, shutil, datetime, zipfile, difflib
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List, Optional, Dict, Any

from PyQt5.QtCore import QThread, pyqtSignal
from PyQt5.QtGui import QColor

from .deps import (
    fitz, PDF2DocxConverter, pytesseract, DocxTemplate,
    PyHankoReader, validate_pdf_signature, pyhanko_signers,
    PdfSigner, PdfSignatureMetadata, PyHankoVC, PYHANKO_OK,
    pypdf, PdfReader, PdfWriter, Transformation,
    Image, ImageOps, Document,
    letter, canvas, rl_colors, SimpleDocTemplate, Paragraph, Spacer,
    getSampleStyleSheet, ImageReader,
    qrcode, pythoncom, win32com
)
from .config import find_tesseract, find_ghostscript
from .helpers import parse_page_ranges, preserve_pdf_metadata_and_structure
from .styles import ACCENT, TEXT_PRIMARY

class ThumbnailWorker(QThread):
    thumb_ready = pyqtSignal(int, bytes)
    finished_all = pyqtSignal(int, str)
    def __init__(self, pdf_path, width=140):
        super().__init__(); self.pdf_path = pdf_path; self.width = width
    def run(self):
        if fitz is None:
            self.finished_all.emit(-1, "PyMuPDF not installed."); return
        try:
            doc = fitz.open(self.pdf_path)
            for i in range(len(doc)):
                rect = doc[i].rect
                if rect.width <= 0: continue
                zoom = self.width / rect.width
                pix = doc[i].get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
                self.thumb_ready.emit(i, pix.tobytes("png"))
            total = len(doc); doc.close()
            self.finished_all.emit(total, "")
        except Exception as ex:
            self.finished_all.emit(-1, str(ex))


class DocWorker(QThread):
    progress = pyqtSignal(int)
    log_line = pyqtSignal(str)
    finished = pyqtSignal(bool, str, str)

    def __init__(self, action, params):
        super().__init__(); self.action = action; self.p = params

    def run(self):
        try:
            dispatch = {
                "split": self._run_split, "split_by_size": self._run_split_by_size,
                "merge": self._run_merge, "convert": self._run_convert,
                "security": self._run_security, "watermark": self._run_watermark,
                "numbering": self._run_numbering, "header_footer": self._run_header_footer,
                "organize": self._run_organize, "insert_pages": self._run_insert_pages,
                "resize": self._run_resize, "remove_blanks": self._run_remove_blanks,
                "metadata": self._run_metadata, "qrcode": self._run_qrcode,
                "style_inspect": self._run_style_inspect, "source_view": self._run_source_view,
                "compress": self._run_compress, "batch": self._run_batch,
                "redact": self._run_redact, "nup": self._run_nup,
                "bates": self._run_bates, "compare": self._run_compare,
                "organize_visual": self._run_organize_visual,
                "ocr": self._run_ocr,
                "verify_signatures": self._run_verify_signatures,
                "sign_pdf": self._run_sign_pdf,
                "pdfa": self._run_pdfa,
                "mail_merge": self._run_mail_merge,
            }
            fn = dispatch.get(self.action)
            if not fn:
                self.finished.emit(False, f"Unknown action: {self.action}", ""); return
            fn()
        except Exception as ex:
            import traceback
            self.finished.emit(False, f"{type(ex).__name__}: {ex}\n{traceback.format_exc()}", "")

    # ── Split ────────────────────────────────────────────────
    def _run_split(self):
        pdf_path = Path(self.p["pdf_path"]); out_dir = Path(self.p["output_dir"])
        mode = self.p["mode"]; reader = PdfReader(str(pdf_path))
        total = len(reader.pages); stem = pdf_path.stem
        out_dir.mkdir(parents=True, exist_ok=True)

        def write_pages(idxs, target):
            w = PdfWriter(); valid = [i for i in idxs if 0 <= i < total]
            w.append(reader, pages=valid, import_outline=True)
            preserve_pdf_metadata_and_structure(reader, w)
            with open(target, "wb") as f: w.write(f)
            self.log_line.emit(f"  Saved: {target.name} ({len(valid)} pages)")

        if mode == "all":
            for i in range(total):
                write_pages([i], out_dir / f"{stem}_page_{i+1:04d}.pdf")
                self.progress.emit(int((i+1)/total*100))
            self.finished.emit(True, f"Extracted {total} pages", str(out_dir))
        elif mode == "single":
            p = self.p["page"] - 1
            out = out_dir / f"{stem}_page_{p+1:04d}.pdf"
            write_pages([p], out); self.progress.emit(100)
            self.finished.emit(True, f"Extracted page {p+1}", str(out))
        elif mode == "range":
            s = self.p["start"] - 1; e = self.p["end"]
            out = out_dir / f"{stem}_pages_{s+1:04d}-{e:04d}.pdf"
            write_pages(list(range(s, e)), out); self.progress.emit(100)
            self.finished.emit(True, f"Extracted {e-s} pages", str(out))
        elif mode == "custom":
            pages = self.p["pages"]
            idxs = [p-1 for p in pages if 0 <= p-1 < total]
            label = "_".join(str(p) for p in pages[:4]) + ("_etc" if len(pages) > 4 else "")
            out = out_dir / f"{stem}_custom_{label}.pdf"
            write_pages(idxs, out); self.progress.emit(100)
            self.finished.emit(True, f"Extracted {len(idxs)} pages", str(out))
        elif mode == "chunk":
            chunk = self.p["chunk"]; chunks = (total + chunk - 1) // chunk
            for i in range(chunks):
                s = i * chunk; e = min(s + chunk, total)
                write_pages(list(range(s, e)), out_dir / f"{stem}_chunk_{i+1:03d}.pdf")
                self.progress.emit(int((i+1)/chunks*100))
            self.finished.emit(True, f"Created {chunks} chunks", str(out_dir))

    def _run_split_by_size(self):
        pdf_path = Path(self.p["pdf_path"]); out_dir = Path(self.p["output_dir"])
        target = int(self.p["target_mb"] * 1024 * 1024)
        out_dir.mkdir(parents=True, exist_ok=True)
        reader = PdfReader(str(pdf_path)); total = len(reader.pages); stem = pdf_path.stem
        part_idx = 1; current = []
        def flush(pages, idx):
            w = PdfWriter(); w.append(reader, pages=pages, import_outline=True)
            preserve_pdf_metadata_and_structure(reader, w)
            out = out_dir / f"{stem}_part_{idx:03d}.pdf"
            with open(out, "wb") as f: w.write(f)
            self.log_line.emit(f"  Part {idx}: {len(pages)} pages ({out.stat().st_size/1024/1024:.2f} MB)")
        for i in range(total):
            current.append(i)
            probe = PdfWriter(); probe.append(reader, pages=current, import_outline=False)
            buf = io.BytesIO(); probe.write(buf)
            if buf.tell() > target and len(current) > 1:
                current.pop(); flush(current, part_idx); part_idx += 1; current = [i]
            self.progress.emit(int((i+1)/total*100))
        if current: flush(current, part_idx)
        self.finished.emit(True, f"Split into {part_idx} parts", str(out_dir))

    # ── Merge ────────────────────────────────────────────────
    def _run_merge(self):
        files = self.p["files"]; out = Path(self.p["out_file"])
        out.parent.mkdir(parents=True, exist_ok=True)
        w = PdfWriter(); n_files = len(files); n_pages = 0
        for i, f in enumerate(files):
            r = PdfReader(f)
            w.append(r, outline_item=Path(f).stem, import_outline=True)
            n_pages += len(r.pages); self.progress.emit(int((i+1)/n_files*85))
        if files:
            try: preserve_pdf_metadata_and_structure(PdfReader(files[0]), w)
            except Exception: pass
        with open(out, "wb") as fh: w.write(fh)
        self.progress.emit(100)
        self.finished.emit(True, f"Merged {n_files} files ({n_pages} pages)", str(out))

    # ── Convert ──────────────────────────────────────────────
    def _run_convert(self):
        ct = self.p["conv_type"]; in_p = self.p["input_path"]
        out = Path(self.p["output_path"]); out.parent.mkdir(parents=True, exist_ok=True)
        if ct == "docx_to_pdf":
            ok = False
            if pythoncom and win32com:
                try:
                    pythoncom.CoInitialize()
                    word = win32com.client.DispatchEx("Word.Application")
                    word.Visible = False; word.DisplayAlerts = 0
                    doc = word.Documents.Open(os.path.abspath(in_p), ReadOnly=True, ConfirmConversions=False)
                    doc.ExportAsFixedFormat(OutputFileName=os.path.abspath(str(out)),
                        ExportFormat=17, OpenAfterExport=False, OptimizeFor=0, CreateBookmarks=1,
                        DocStructureTags=True, BitmapMissingFonts=True, IncludeDocProps=True)
                    doc.Close(False); word.Quit(); pythoncom.CoUninitialize(); ok = True
                except Exception as ex: self.log_line.emit(f"  Word COM: {ex}")
            if not ok:
                from docx2pdf import convert; convert(in_p, str(out))
            self.progress.emit(100); self.finished.emit(True, "Word ➔ PDF", str(out))
        elif ct == "excel_to_pdf":
            if not (pythoncom and win32com): raise RuntimeError("Excel COM requires pywin32")
            pythoncom.CoInitialize()
            xl = win32com.client.DispatchEx("Excel.Application")
            xl.Visible = False; xl.DisplayAlerts = False
            wb = xl.Workbooks.Open(os.path.abspath(in_p), ReadOnly=True)
            wb.ExportAsFixedFormat(0, os.path.abspath(str(out)), 0, True, False)
            wb.Close(False); xl.Quit(); pythoncom.CoUninitialize()
            self.progress.emit(100); self.finished.emit(True, "Excel ➔ PDF", str(out))
        elif ct == "images_to_pdf":
            imgs = []
            for p in self.p["img_paths"]:
                im = Image.open(p)
                if ImageOps: im = ImageOps.exif_transpose(im)
                if im.mode not in ("RGB", "L"): im = im.convert("RGB")
                imgs.append(im)
            imgs[0].save(str(out), save_all=True, append_images=imgs[1:], resolution=100.0)
            self.progress.emit(100); self.finished.emit(True, f"Combined {len(imgs)} images", str(out))
        elif ct == "pdf_to_images":
            reader = PdfReader(in_p)
            out_dir = out if out.is_dir() else out.parent / (Path(in_p).stem + "_images")
            out_dir.mkdir(parents=True, exist_ok=True); cnt = 0
            for i, page in enumerate(reader.pages):
                for j, img in enumerate(page.images):
                    with open(out_dir / f"page_{i+1:03d}_img_{j+1}_{img.name}", "wb") as f: f.write(img.data)
                    cnt += 1
                self.progress.emit(int((i+1)/len(reader.pages)*100))
            self.finished.emit(True, f"Extracted {cnt} images", str(out_dir))
        elif ct == "pdf_to_docx":
            out_str = str(out)
            converted = False
            if PDF2DocxConverter:
                try:
                    self.log_line.emit("  Preserving formatting: analyzing layout, fonts, colors, and tables...")
                    import logging
                    pat = re.compile(r'\((\d+)/(\d+)\)')
                    worker_self = self

                    class _ConvProgressHandler(logging.Handler):
                        def emit(self, record):
                            try:
                                msg = record.getMessage()
                                m = pat.search(msg)
                                if m:
                                    cur, tot = int(m.group(1)), int(m.group(2))
                                    if tot > 0:
                                        worker_self.progress.emit(min(95, max(5, int(cur / tot * 90))))
                            except Exception: pass

                    handler = _ConvProgressHandler()
                    p2d_logger = logging.getLogger("pdf2docx")
                    p2d_logger.addHandler(handler)
                    try:
                        cv = PDF2DocxConverter(in_p)
                        cv.convert(out_str)
                        cv.close()
                        converted = True
                    finally:
                        p2d_logger.removeHandler(handler)
                except Exception as ex:
                    self.log_line.emit(f"  pdf2docx note: {ex}. Falling back to standard document extractor...")

            if not converted:
                if not Document: raise RuntimeError("python-docx not installed")
                reader = PdfReader(in_p); doc = Document()
                if reader.metadata:
                    for src, attr in [("/Title","title"),("/Author","author"),("/Subject","subject")]:
                        v = reader.metadata.get(src)
                        if v: setattr(doc.core_properties, attr, str(v))
                total = len(reader.pages)
                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        for para in text.split("\n\n"):
                            c = para.strip()
                            if not c: continue
                            if len(c) < 60 and c.isupper() and not c.endswith("."): doc.add_heading(c, 2)
                            elif c.startswith(("• ","- ","* ")): doc.add_paragraph(c[2:], style="List Bullet")
                            else: doc.add_paragraph(c)
                    if i < total - 1: doc.add_page_break()
                    self.progress.emit(int((i+1)/total*100))
                doc.save(out_str)

            self.progress.emit(100)
            status_msg = "PDF ➔ Word (Layout Preserved)" if converted else "PDF ➔ Word"
            self.finished.emit(True, status_msg, out_str)
        elif ct == "pdf_to_txt":
            reader = PdfReader(in_p); total = len(reader.pages)
            with open(out, "w", encoding="utf-8") as f:
                for i, page in enumerate(reader.pages):
                    f.write(f"--- PAGE {i+1} ---\n\n"); f.write((page.extract_text() or "") + "\n\n")
                    self.progress.emit(int((i+1)/total*100))
            self.finished.emit(True, "PDF ➔ Text", str(out))
        elif ct == "text_to_pdf":
            if not SimpleDocTemplate: raise RuntimeError("reportlab not installed")
            doc = SimpleDocTemplate(str(out), pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
            styles = getSampleStyleSheet(); story = []
            with open(in_p, "r", encoding="utf-8", errors="replace") as f:
                for line in f.readlines():
                    c = line.rstrip("\r\n")
                    if c.startswith("# "): story.append(Paragraph(c[2:], styles["Heading1"]))
                    elif c.startswith("## "): story.append(Paragraph(c[3:], styles["Heading2"]))
                    elif c.startswith("### "): story.append(Paragraph(c[4:], styles["Heading3"]))
                    elif not c.strip(): story.append(Spacer(1, 10))
                    else:
                        safe = c.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
                        story.append(Paragraph(safe, styles["Normal"]))
            doc.build(story); self.progress.emit(100)
            self.finished.emit(True, "Text ➔ PDF", str(out))

    # ── Security ─────────────────────────────────────────────
    def _run_security(self):
        mode = self.p["sec_mode"]; in_p = self.p["pdf_path"]
        out = Path(self.p["out_file"]); out.parent.mkdir(parents=True, exist_ok=True)
        if mode == "protect":
            r = PdfReader(in_p); w = PdfWriter()
            w.append(r, import_outline=True)
            preserve_pdf_metadata_and_structure(r, w)
            perms = None
            if hasattr(pypdf.constants, "UserAccessPermissions"):
                flag = 0
                if self.p.get("allow_print", True): flag |= pypdf.constants.UserAccessPermissions.PRINT
                if self.p.get("allow_copy", True):  flag |= pypdf.constants.UserAccessPermissions.EXTRACT
                perms = flag
            w.encrypt(user_password=self.p["user_pw"],
                      owner_password=self.p["owner_pw"] or self.p["user_pw"],
                      permissions_flag=perms)
            with open(out, "wb") as f: w.write(f)
            self.progress.emit(100); self.finished.emit(True, "PDF encrypted", str(out))
        elif mode == "unlock":
            r = PdfReader(in_p)
            if r.is_encrypted and r.decrypt(self.p["password"]) == 0:
                raise ValueError("Incorrect password.")
            w = PdfWriter(); w.append(r, import_outline=True)
            preserve_pdf_metadata_and_structure(r, w)
            with open(out, "wb") as f: w.write(f)
            self.progress.emit(100); self.finished.emit(True, "Password removed", str(out))

    # ── Watermark / Numbering / Header-Footer ────────────────
    def _run_watermark(self):
        r = PdfReader(self.p["pdf_path"]); w = PdfWriter()
        w.append(r, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)
        total = len(w.pages)
        wm_type = self.p.get("wm_type", "text")
        is_bg = self.p.get("is_background", True)

        if wm_type == "image":
            img_path = self.p.get("image_path")
            if not img_path or not os.path.exists(img_path):
                raise ValueError("Watermark image not found or not specified.")
            if not Image:
                raise RuntimeError("Pillow (PIL) is required for image watermarks.")
            if not ImageReader:
                raise RuntimeError("ReportLab ImageReader is required for image watermarks.")

            angle = float(self.p.get("angle", 0))
            op = max(0.01, min(1.0, float(self.p.get("opacity", 0.3))))
            scale_pct = max(5.0, min(150.0, float(self.p.get("scale_pct", 50))))
            pos = self.p.get("position", "center")

            raw_img = Image.open(img_path).convert("RGBA")
            if op < 1.0:
                r_ch, g_ch, b_ch, a_ch = raw_img.split()
                a_ch = a_ch.point(lambda p: int(p * op))
                raw_img.putalpha(a_ch)
            img_reader = ImageReader(raw_img)
            orig_w, orig_h = raw_img.size
            aspect = orig_h / orig_w if orig_w else 1.0

            for i, page in enumerate(w.pages):
                pw = float(page.mediabox.width); ph = float(page.mediabox.height)
                target_w = pw * (scale_pct / 100.0)
                target_h = target_w * aspect
                if target_h > ph * (scale_pct / 100.0):
                    target_h = ph * (scale_pct / 100.0)
                    target_w = target_h / aspect

                buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
                c.saveState()

                if pos == "tiled":
                    step_x = max(60, target_w * 1.4)
                    step_y = max(60, target_h * 1.4)
                    y = step_y / 2
                    while y < ph + step_y:
                        x = step_x / 2
                        while x < pw + step_x:
                            c.saveState()
                            c.translate(x, y); c.rotate(angle)
                            c.drawImage(img_reader, -target_w / 2, -target_h / 2, width=target_w, height=target_h, mask='auto')
                            c.restoreState()
                            x += step_x
                        y += step_y
                else:
                    m = 36
                    if pos == "top-left": cx, cy = target_w / 2 + m, ph - target_h / 2 - m
                    elif pos == "top-right": cx, cy = pw - target_w / 2 - m, ph - target_h / 2 - m
                    elif pos == "bottom-left": cx, cy = target_w / 2 + m, target_h / 2 + m
                    elif pos == "bottom-right": cx, cy = pw - target_w / 2 - m, target_h / 2 + m
                    else: cx, cy = pw / 2, ph / 2

                    c.translate(cx, cy); c.rotate(angle)
                    c.drawImage(img_reader, -target_w / 2, -target_h / 2, width=target_w, height=target_h, mask='auto')

                c.restoreState(); c.showPage(); c.save(); buf.seek(0)
                page.merge_page(PdfReader(buf).pages[0], over=not is_bg)
                self.progress.emit(int((i+1)/total*100))

            with open(self.p["out_file"], "wb") as f: w.write(f)
            self.finished.emit(True, "Image watermark applied", self.p["out_file"])
        else:
            qcol = QColor(self.p.get("color_hex", ACCENT))
            rr, gg, bb = qcol.red()/255, qcol.green()/255, qcol.blue()/255
            for i, page in enumerate(w.pages):
                pw = float(page.mediabox.width); ph = float(page.mediabox.height)
                buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
                c.saveState(); c.translate(pw/2, ph/2); c.rotate(self.p.get("angle", 45))
                c.setFillColor(rl_colors.Color(rr, gg, bb, alpha=self.p.get("opacity", 0.25)))
                c.setFont("Helvetica-Bold", self.p.get("font_size", 56))
                c.drawCentredString(0, 0, self.p.get("text", "CONFIDENTIAL"))
                c.restoreState(); c.showPage(); c.save(); buf.seek(0)
                page.merge_page(PdfReader(buf).pages[0], over=not is_bg)
                self.progress.emit(int((i+1)/total*100))
            with open(self.p["out_file"], "wb") as f: w.write(f)
            self.finished.emit(True, "Watermark applied", self.p["out_file"])

    def _run_numbering(self):
        r = PdfReader(self.p["pdf_path"]); w = PdfWriter()
        w.append(r, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)
        total = len(w.pages); fmt = self.p["format"]; pos = self.p["position"]
        for i, page in enumerate(w.pages):
            pw = float(page.mediabox.width); ph = float(page.mediabox.height); cur = i + 1
            label = {"Page X of Y": f"Page {cur} of {total}", "Page X": f"Page {cur}",
                     "X / Y": f"{cur} / {total}", "- X -": f"- {cur} -"}.get(fmt, str(cur))
            buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
            c.setFont("Helvetica", self.p["font_size"]); c.setFillColor(rl_colors.HexColor(TEXT_PRIMARY))
            m = 36
            if pos == "bottom-center": c.drawCentredString(pw/2, m, label)
            elif pos == "bottom-right": c.drawRightString(pw-m, m, label)
            elif pos == "bottom-left": c.drawString(m, m, label)
            elif pos == "top-center": c.drawCentredString(pw/2, ph-m, label)
            elif pos == "top-right": c.drawRightString(pw-m, ph-m, label)
            c.showPage(); c.save(); buf.seek(0)
            page.merge_page(PdfReader(buf).pages[0], over=True)
            self.progress.emit(int((i+1)/total*100))
        with open(self.p["out_file"], "wb") as f: w.write(f)
        self.finished.emit(True, "Page numbering added", self.p["out_file"])

    def _run_header_footer(self):
        r = PdfReader(self.p["pdf_path"]); w = PdfWriter()
        w.append(r, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)
        total = len(w.pages); header = self.p.get("header_text", "")
        footer = self.p.get("footer_text", ""); fsz = int(self.p.get("font_size", 9))
        fname = Path(self.p["pdf_path"]).name; now = datetime.datetime.now()
        title = ""
        try: title = (r.metadata or {}).get("/Title") or ""
        except Exception: pass
        qcol = QColor(self.p.get("color_hex", TEXT_PRIMARY))
        rr, gg, bb = qcol.red()/255, qcol.green()/255, qcol.blue()/255
        def fmt(t, cur):
            return (t.replace("{page}", str(cur)).replace("{pages}", str(total))
                     .replace("{filename}", fname).replace("{title}", str(title))
                     .replace("{date}", now.strftime("%Y-%m-%d")).replace("{time}", now.strftime("%H:%M")))
        for i, page in enumerate(w.pages):
            pw = float(page.mediabox.width); ph = float(page.mediabox.height); cur = i + 1
            if not (header.strip() or footer.strip()): continue
            buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
            c.setFont("Helvetica", fsz); c.setFillColor(rl_colors.Color(rr, gg, bb, alpha=1))
            if header.strip(): c.drawCentredString(pw/2, ph-32, fmt(header, cur))
            if footer.strip(): c.drawCentredString(pw/2, 32, fmt(footer, cur))
            c.showPage(); c.save(); buf.seek(0)
            page.merge_page(PdfReader(buf).pages[0], over=True)
            self.progress.emit(int((i+1)/total*100))
        with open(self.p["out_file"], "wb") as f: w.write(f)
        self.finished.emit(True, "Header/footer added", self.p["out_file"])

    # ── Organize ─────────────────────────────────────────────
    def _run_organize(self):
        mode = self.p["org_mode"]; r = PdfReader(self.p["pdf_path"])
        total = len(r.pages); out = Path(self.p["out_file"])
        if mode == "rotate":
            w = PdfWriter(); w.append(r, import_outline=True)
            preserve_pdf_metadata_and_structure(r, w)
            sel = set(self.p.get("pages", []))
            for i, page in enumerate(w.pages):
                if self.p["scope"] == "all" or (i+1) in sel: page.rotate(self.p["angle"])
                self.progress.emit(int((i+1)/total*100))
            with open(out, "wb") as f: w.write(f)
            self.finished.emit(True, f"Rotated {self.p['angle']}°", str(out))
        elif mode == "delete":
            del_set = set(self.p["delete_pages"])
            kept = [i for i in range(total) if (i+1) not in del_set]
            w = PdfWriter(); w.append(r, pages=kept, import_outline=True)
            preserve_pdf_metadata_and_structure(r, w); self.progress.emit(100)
            with open(out, "wb") as f: w.write(f)
            self.finished.emit(True, f"Removed {len(del_set)} pages", str(out))
        elif mode == "reverse":
            w = PdfWriter(); w.append(r, pages=list(reversed(range(total))), import_outline=False)
            preserve_pdf_metadata_and_structure(r, w); self.progress.emit(100)
            with open(out, "wb") as f: w.write(f)
            self.finished.emit(True, "Reversed page order", str(out))

    def _run_insert_pages(self):
        base = PdfReader(self.p["base_pdf"]); ins = PdfReader(self.p["insert_pdf"])
        out = Path(self.p["out_file"]); out.parent.mkdir(parents=True, exist_ok=True)
        pos = max(1, min(int(self.p["position"]), len(base.pages) + 1))
        pages = self.p.get("pages") or None
        idxs = [p-1 for p in pages if 0 <= p-1 < len(ins.pages)] if pages else list(range(len(ins.pages)))
        w = PdfWriter()
        for i in range(pos - 1): w.add_page(base.pages[i])
        for i in idxs: w.add_page(ins.pages[i])
        for i in range(pos - 1, len(base.pages)): w.add_page(base.pages[i])
        preserve_pdf_metadata_and_structure(base, w)
        with open(out, "wb") as f: w.write(f)
        self.progress.emit(100)
        self.finished.emit(True, f"Inserted {len(idxs)} pages at position {pos}", str(out))

    def _run_resize(self):
        r = PdfReader(self.p["pdf_path"]); w = PdfWriter()
        out = Path(self.p["out_file"]); out.parent.mkdir(parents=True, exist_ok=True)
        PT_MM = 72.0 / 25.4
        SIZES = {"a4": (595.276, 841.890), "letter": (612.0, 792.0),
                 "legal": (612.0, 1008.0), "a5": (419.528, 595.276)}
        target = self.p.get("target", "a4"); margin_mm = float(self.p.get("margin_mm", 0))
        total = len(r.pages)
        for i, page in enumerate(r.pages):
            sw = float(page.mediabox.width); sh = float(page.mediabox.height)
            if target == "fit":
                tw, th = SIZES["a4"]; scale = (tw - 2*margin_mm*PT_MM) / sw
                th = sh * scale + 2*margin_mm*PT_MM
            else:
                tw, th = SIZES.get(target, SIZES["a4"])
            inner_w = tw - 2*margin_mm*PT_MM; inner_h = th - 2*margin_mm*PT_MM
            scale = min(inner_w / sw, inner_h / sh)
            tx = margin_mm*PT_MM + (inner_w - sw*scale) / 2.0
            ty = margin_mm*PT_MM + (inner_h - sh*scale) / 2.0
            np_ = w.add_blank_page(width=tw, height=th)
            np_.merge_transformed_page(page, Transformation().scale(scale, scale).translate(tx, ty), over=True)
            self.progress.emit(int((i+1)/total*100))
        preserve_pdf_metadata_and_structure(r, w)
        with open(out, "wb") as f: w.write(f)
        self.finished.emit(True, f"Resized {total} pages to {target.upper()}", str(out))

    def _run_remove_blanks(self):
        r = PdfReader(self.p["pdf_path"]); total = len(r.pages)
        th = float(self.p.get("threshold", 0.995))
        out = Path(self.p["out_file"]); out.parent.mkdir(parents=True, exist_ok=True)
        keep = []; fitz_used = False
        if fitz is not None:
            try:
                doc = fitz.open(self.p["pdf_path"])
                for i in range(len(doc)):
                    pix = doc[i].get_pixmap(dpi=40, colorspace=fitz.csGRAY)
                    s = pix.samples
                    if not s: keep.append(i); continue
                    dark = sum(1 for v in s if v < 200)
                    if dark / len(s) > (1.0 - th): keep.append(i)
                    self.progress.emit(int((i+1)/total*60))
                doc.close(); fitz_used = True
            except Exception: pass
        if not fitz_used:
            for i, page in enumerate(r.pages):
                try:
                    text = (page.extract_text() or "").strip()
                    c = page.get_contents(); raw = b""
                    if c is not None:
                        if hasattr(c, "get_data"): raw = c.get_data()
                        else:
                            for item in c:
                                try: raw += item.get_data()
                                except Exception: pass
                    if text or len(raw) >= 200: keep.append(i)
                except Exception: keep.append(i)
                self.progress.emit(int((i+1)/total*60))
        removed = total - len(keep)
        if removed == 0:
            self.finished.emit(True, "No blank pages detected.", self.p["pdf_path"]); return
        w = PdfWriter(); w.append(r, pages=keep, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)
        with open(out, "wb") as f: w.write(f)
        self.progress.emit(100)
        self.finished.emit(True, f"Removed {removed} blank pages; {len(keep)} kept", str(out))

    def _run_metadata(self):
        r = PdfReader(self.p["pdf_path"]); w = PdfWriter()
        w.append(r, import_outline=True)
        if self.p.get("clear_all"):
            w.add_metadata({})
        else:
            meta = {}
            if r.metadata:
                for k, v in r.metadata.items():
                    if v is not None: meta[k] = str(v)
            for key in ("title", "author", "subject", "keywords", "creator"):
                if self.p.get(key): meta["/" + key.capitalize()] = self.p[key]
            w.add_metadata(meta)
        with open(self.p["out_file"], "wb") as f: w.write(f)
        self.progress.emit(100); self.finished.emit(True, "Metadata saved", self.p["out_file"])

    def _run_qrcode(self):
        if not qrcode: raise RuntimeError("qrcode not installed.")
        qr = qrcode.QRCode(version=1, box_size=self.p.get("box_size", 10), border=self.p.get("border", 4))
        qr.add_data(self.p["text"]); qr.make(fit=True)
        qr.make_image(fill_color="black", back_color="white").save(self.p["out_file"])
        self.progress.emit(100); self.finished.emit(True, "QR Code generated", self.p["out_file"])

    def _run_compress(self):
        in_p = self.p["pdf_path"]; out = Path(self.p["out_file"])
        preset = self.p.get("preset", "lossless")
        out.parent.mkdir(parents=True, exist_ok=True)
        size_before = Path(in_p).stat().st_size
        r = PdfReader(in_p); w = PdfWriter()
        w.append(r, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)
        for i, page in enumerate(w.pages):
            try: page.compress_content_streams()
            except Exception: pass
            self.progress.emit(10 + int((i+1)/len(w.pages)*30))
        try: w.compress_identical_objects(remove_identicals=True, remove_orphans=True)
        except Exception: pass
        tmp = out.with_suffix(".lossless.tmp.pdf")
        with open(tmp, "wb") as f: w.write(f)
        self.progress.emit(50)
        note = " (lossless)"
        if preset in ("balanced", "aggressive") and fitz is not None:
            try:
                dpi_cap = 150 if preset == "balanced" else 110
                jq = 78 if preset == "balanced" else 62
                doc = fitz.open(str(tmp))
                for pno in range(len(doc)):
                    for img in doc[pno].get_images(full=True):
                        try:
                            base = doc.extract_image(img[0])
                            if base["ext"] not in ("png", "jpeg", "jpg"): continue
                            im = Image.open(io.BytesIO(base["image"]))
                            w_, h_ = im.size
                            if w_ * h_ < 40000: continue
                            mx = int(dpi_cap * 8.5)
                            if w_ > mx:
                                r_ = mx / w_; im = im.resize((mx, int(h_*r_)), Image.LANCZOS)
                            buf = io.BytesIO()
                            if im.mode in ("RGBA", "P"): im = im.convert("RGB")
                            im.save(buf, format="JPEG", quality=jq, optimize=True)
                            doc.update_stream(img[0], buf.getvalue(), compress=True)
                        except Exception: continue
                    self.progress.emit(50 + int((pno+1)/len(doc)*40))
                doc.save(str(out), garbage=4, deflate=True, clean=True); doc.close()
                tmp.unlink(missing_ok=True)
                note = f" (images @ {dpi_cap} DPI, q{jq})"
            except Exception as ex:
                self.log_line.emit(f"  Image pass skipped: {ex}")
                shutil.move(str(tmp), str(out))
        else:
            shutil.move(str(tmp), str(out))
            if preset != "lossless" and fitz is None:
                self.log_line.emit("  Install PyMuPDF for deeper compression.")
        sa = out.stat().st_size; saved = size_before - sa
        pct = (saved / size_before * 100) if size_before else 0
        self.progress.emit(100)
        self.finished.emit(True, f"Compressed {size_before/1024:.0f} KB ➔ {sa/1024:.0f} KB ({pct:.0f}% smaller){note}", str(out))

    # ── Batch ────────────────────────────────────────────────
    def _run_batch(self):
        in_d = Path(self.p["input_dir"]); out_d = Path(self.p["output_dir"])
        op = self.p["batch_op"]; pat = self.p.get("pattern", "*.pdf")
        out_d.mkdir(parents=True, exist_ok=True)
        files = [f for f in sorted(in_d.glob(pat)) if f.is_file() and f.suffix.lower() == ".pdf"]
        if not files:
            self.finished.emit(False, f"No PDFs matched {pat}", ""); return
        ok = 0; fail = []
        for idx, f in enumerate(files):
            self.log_line.emit(f"[{idx+1}/{len(files)}] {f.name}")
            out = out_d / f.name
            try:
                if op == "compress": self._batch_compress(f, out)
                elif op == "watermark": self._batch_watermark(f, out)
                elif op == "numbering": self._batch_numbering(f, out)
                elif op == "metadata_clear": self._batch_metadata_clear(f, out)
                elif op == "remove_blanks": self._batch_remove_blanks(f, out)
                ok += 1
            except Exception as ex:
                self.log_line.emit(f"  FAILED: {ex}"); fail.append(f.name)
            self.progress.emit(int((idx+1)/len(files)*100))
        msg = f"Batch: {ok} succeeded, {len(fail)} failed"
        self.finished.emit(ok > 0, msg, str(out_d))

    def _batch_compress(self, src, dst):
        r = PdfReader(str(src)); w = PdfWriter()
        w.append(r, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)
        for p in w.pages:
            try: p.compress_content_streams()
            except Exception: pass
        try: w.compress_identical_objects(remove_identicals=True, remove_orphans=True)
        except Exception: pass
        with open(dst, "wb") as f: w.write(f)
        self.log_line.emit(f"  -> {dst.name} ({dst.stat().st_size/1024:.0f} KB)")

    def _batch_watermark(self, src, dst):
        wm_type = self.p.get("wm_type", "text")
        is_bg = self.p.get("is_background", True)
        r = PdfReader(str(src)); w = PdfWriter()
        w.append(r, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)

        if wm_type == "image" and self.p.get("image_path") and os.path.exists(self.p.get("image_path")):
            img_path = self.p["image_path"]
            angle = float(self.p.get("angle", 0))
            op = max(0.01, min(1.0, float(self.p.get("opacity", 0.3))))
            scale_pct = max(5.0, min(150.0, float(self.p.get("scale_pct", 50))))
            pos = self.p.get("position", "center")

            raw_img = Image.open(img_path).convert("RGBA")
            if op < 1.0:
                r_ch, g_ch, b_ch, a_ch = raw_img.split()
                a_ch = a_ch.point(lambda p: int(p * op))
                raw_img.putalpha(a_ch)
            img_reader = ImageReader(raw_img)
            orig_w, orig_h = raw_img.size
            aspect = orig_h / orig_w if orig_w else 1.0

            for page in w.pages:
                pw = float(page.mediabox.width); ph = float(page.mediabox.height)
                target_w = pw * (scale_pct / 100.0)
                target_h = target_w * aspect
                if target_h > ph * (scale_pct / 100.0):
                    target_h = ph * (scale_pct / 100.0)
                    target_w = target_h / aspect

                buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
                c.saveState()
                if pos == "tiled":
                    step_x = max(60, target_w * 1.4)
                    step_y = max(60, target_h * 1.4)
                    y = step_y / 2
                    while y < ph + step_y:
                        x = step_x / 2
                        while x < pw + step_x:
                            c.saveState(); c.translate(x, y); c.rotate(angle)
                            c.drawImage(img_reader, -target_w / 2, -target_h / 2, width=target_w, height=target_h, mask='auto')
                            c.restoreState()
                            x += step_x
                        y += step_y
                else:
                    m = 36
                    if pos == "top-left": cx, cy = target_w / 2 + m, ph - target_h / 2 - m
                    elif pos == "top-right": cx, cy = pw - target_w / 2 - m, ph - target_h / 2 - m
                    elif pos == "bottom-left": cx, cy = target_w / 2 + m, target_h / 2 + m
                    elif pos == "bottom-right": cx, cy = pw - target_w / 2 - m, target_h / 2 + m
                    else: cx, cy = pw / 2, ph / 2
                    c.translate(cx, cy); c.rotate(angle)
                    c.drawImage(img_reader, -target_w / 2, -target_h / 2, width=target_w, height=target_h, mask='auto')
                c.restoreState(); c.showPage(); c.save(); buf.seek(0)
                page.merge_page(PdfReader(buf).pages[0], over=not is_bg)
        else:
            text = self.p.get("text", "CONFIDENTIAL"); angle = self.p.get("angle", 45)
            op = self.p.get("opacity", 0.25); fsz = self.p.get("font_size", 56)
            qcol = QColor(self.p.get("color_hex", ACCENT))
            rr, gg, bb = qcol.red()/255, qcol.green()/255, qcol.blue()/255
            for page in w.pages:
                pw = float(page.mediabox.width); ph = float(page.mediabox.height)
                buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
                c.saveState(); c.translate(pw/2, ph/2); c.rotate(angle)
                c.setFillColor(rl_colors.Color(rr, gg, bb, alpha=op))
                c.setFont("Helvetica-Bold", fsz); c.drawCentredString(0, 0, text)
                c.restoreState(); c.showPage(); c.save(); buf.seek(0)
                page.merge_page(PdfReader(buf).pages[0], over=not is_bg)

        with open(dst, "wb") as f: w.write(f)
        self.log_line.emit(f"  -> {dst.name}")

    def _batch_numbering(self, src, dst):
        fmt = self.p.get("format", "Page X of Y"); pos = self.p.get("position", "bottom-center")
        fsz = self.p.get("font_size", 10)
        r = PdfReader(str(src)); w = PdfWriter()
        w.append(r, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)
        total = len(w.pages)
        for i, page in enumerate(w.pages):
            pw = float(page.mediabox.width); ph = float(page.mediabox.height); cur = i + 1
            label = {"Page X of Y": f"Page {cur} of {total}", "Page X": f"Page {cur}",
                     "X / Y": f"{cur} / {total}", "- X -": f"- {cur} -", "X": str(cur)}.get(fmt, str(cur))
            buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
            c.setFont("Helvetica", fsz); c.setFillColor(rl_colors.HexColor("#0f172a"))
            m = 36
            if pos == "bottom-center": c.drawCentredString(pw/2, m, label)
            elif pos == "bottom-right": c.drawRightString(pw-m, m, label)
            elif pos == "bottom-left": c.drawString(m, m, label)
            elif pos == "top-center": c.drawCentredString(pw/2, ph-m, label)
            elif pos == "top-right": c.drawRightString(pw-m, ph-m, label)
            c.showPage(); c.save(); buf.seek(0)
            page.merge_page(PdfReader(buf).pages[0], over=True)
        with open(dst, "wb") as f: w.write(f)
        self.log_line.emit(f"  -> {dst.name}")

    def _batch_metadata_clear(self, src, dst):
        r = PdfReader(str(src)); w = PdfWriter()
        w.append(r, import_outline=True); w.add_metadata({})
        with open(dst, "wb") as f: w.write(f)
        self.log_line.emit(f"  -> {dst.name}")

    def _batch_remove_blanks(self, src, dst):
        r = PdfReader(str(src)); total = len(r.pages); keep = []
        for i, page in enumerate(r.pages):
            try:
                text = (page.extract_text() or "").strip()
                if text: keep.append(i); continue
                c = page.get_contents(); raw = b""
                if c is not None:
                    if hasattr(c, "get_data"): raw = c.get_data()
                    else:
                        for item in c:
                            try: raw += item.get_data()
                            except Exception: pass
                if len(raw) >= 200: keep.append(i)
            except Exception: keep.append(i)
        if not keep: keep = [0]
        w = PdfWriter(); w.append(r, pages=keep, import_outline=True)
        preserve_pdf_metadata_and_structure(r, w)
        with open(dst, "wb") as f: w.write(f)
        self.log_line.emit(f"  -> {dst.name} ({len(keep)}/{total})")

    # ── Redact ───────────────────────────────────────────────
    def _run_redact(self):
        if fitz is None: raise RuntimeError("PyMuPDF required.")
        in_p = self.p["pdf_path"]; out = Path(self.p["out_file"])
        out.parent.mkdir(parents=True, exist_ok=True)
        terms = [t.strip() for t in self.p.get("terms", "").split(",") if t.strip()]
        if not terms: raise RuntimeError("Provide at least one term.")
        spec = self.p.get("page_range", "").strip()
        fill = (0, 0, 0) if self.p.get("fill", "black") == "black" else (1, 1, 1)
        doc = fitz.open(in_p)
        try:
            if spec:
                rng = parse_page_ranges(spec, len(doc))
                tpg = [p-1 for p in rng if 0 <= p-1 < len(doc)]
            else:
                tpg = list(range(len(doc)))
            hits = 0
            for idx, pno in enumerate(tpg):
                page = doc[pno]; ph = 0
                for term in terms:
                    try: found = page.search_for(term)
                    except Exception: found = []
                    for rect in found:
                        page.add_redact_annot(fitz.Rect(rect.x0-1, rect.y0-1, rect.x1+1, rect.y1+1), fill=fill)
                        ph += 1
                if ph:
                    try: page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS)
                    except TypeError: page.apply_redactions()
                    hits += ph
                    self.log_line.emit(f"  Page {pno+1}: {ph} redacted")
                self.progress.emit(int((idx+1)/len(tpg)*95))
            if hits == 0:
                doc.close(); self.finished.emit(True, "No matches found.", in_p); return
            doc.save(str(out), garbage=4, deflate=True, clean=True); doc.close()
            self.progress.emit(100)
            self.finished.emit(True, f"Redacted {hits} occurrence(s)", str(out))
        except Exception:
            try: doc.close()
            except Exception: pass
            raise

    # ── N-up / Booklet ───────────────────────────────────────
    def _run_nup(self):
        in_p = self.p["pdf_path"]; out = Path(self.p["out_file"])
        out.parent.mkdir(parents=True, exist_ok=True)
        layout = self.p.get("layout", "2up")
        sheet = self.p.get("sheet_size", "a4")
        gutter_mm = float(self.p.get("gutter_mm", 5))
        PT_MM = 72.0 / 25.4
        S = {"a4": (595.276, 841.890), "letter": (612.0, 792.0), "legal": (612.0, 1008.0)}
        sw, sh = S.get(sheet, S["a4"])
        r = PdfReader(in_p); total = len(r.pages)
        if total == 0: self.finished.emit(False, "Empty PDF", ""); return
        cfg = {"2up": (1, 2, "portrait"), "4up": (2, 2, "portrait"),
               "9up": (3, 3, "portrait"), "booklet": (2, 1, "landscape")}
        rows, cols, orient = cfg.get(layout, (1, 2, "portrait"))
        if orient == "landscape": pw_, ph_ = sh, sw
        else: pw_, ph_ = sw, sh
        g = gutter_mm * PT_MM
        cw = max((pw_ - g * (cols + 1)) / cols, 10)
        ch = max((ph_ - g * (rows + 1)) / rows, 10)
        if layout == "booklet":
            padded = list(range(total))
            while len(padded) % 4 != 0: padded.append(None)
            order = []; L = 0; R = len(padded) - 1
            while L < R:
                order.append(padded[R]); order.append(padded[L])
                L += 1; R -= 1
                if L < R:
                    order.append(padded[L]); order.append(padded[R])
                    L += 1; R -= 1
        else:
            order = list(range(total))
        per = rows * cols
        w = PdfWriter(); sheets = (len(order) + per - 1) // per
        for s in range(sheets):
            np_ = w.add_blank_page(width=pw_, height=ph_)
            for slot in range(per):
                i = s * per + slot
                if i >= len(order): break
                sp = order[i]
                if sp is None or sp >= total: continue
                src = r.pages[sp]
                rr_ = slot // cols; cc = slot % cols
                tx = g + cc * (cw + g)
                ty = ph_ - (g + (rr_ + 1) * ch + rr_ * g)
                w0 = float(src.mediabox.width); h0 = float(src.mediabox.height)
                if w0 <= 0 or h0 <= 0: continue
                scale = min(cw / w0, ch / h0)
                ox = tx + (cw - w0 * scale) / 2.0
                oy = ty + (ch - h0 * scale) / 2.0
                np_.merge_transformed_page(src, Transformation().scale(scale, scale).translate(ox, oy), over=True)
            self.progress.emit(int((s+1)/sheets*100))
        preserve_pdf_metadata_and_structure(r, w)
        with open(out, "wb") as f: w.write(f)
        self.finished.emit(True, f"Created {sheets} imposed sheet(s) — {layout.upper()}", str(out))

    # ── Bates ────────────────────────────────────────────────
    def _run_bates(self):
        files = self.p["files"]; out = Path(self.p["out_file"])
        out.parent.mkdir(parents=True, exist_ok=True)
        prefix = self.p.get("prefix", ""); suffix = self.p.get("suffix", "")
        start = int(self.p.get("start", 1)); digits = int(self.p.get("digits", 6))
        fsz = int(self.p.get("font_size", 9)); pos = self.p.get("position", "bottom-right")
        qcol = QColor(self.p.get("color_hex", "#000000"))
        rr, gg, bb = qcol.red()/255, qcol.green()/255, qcol.blue()/255
        margin = 36
        def stamp(num, page):
            pw = float(page.mediabox.width); ph = float(page.mediabox.height)
            buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
            c.setFont("Helvetica", fsz); c.setFillColor(rl_colors.Color(rr, gg, bb, alpha=1))
            if pos == "bottom-right": c.drawRightString(pw-margin, margin, num)
            elif pos == "bottom-left": c.drawString(margin, margin, num)
            elif pos == "bottom-center": c.drawCentredString(pw/2, margin, num)
            elif pos == "top-right": c.drawRightString(pw-margin, ph-margin, num)
            elif pos == "top-left": c.drawString(margin, ph-margin, num)
            c.showPage(); c.save(); buf.seek(0)
            page.merge_page(PdfReader(buf).pages[0], over=True)
        merge = self.p.get("merge_output", True)
        if merge:
            w = PdfWriter(); counter = start; n = len(files)
            for fi, f in enumerate(files):
                r = PdfReader(f)
                for page in r.pages:
                    stamp(f"{prefix}{str(counter).zfill(digits)}{suffix}", page)
                    w.add_page(page); counter += 1
                self.progress.emit(int((fi+1)/n*95))
            if files:
                try: preserve_pdf_metadata_and_structure(PdfReader(files[0]), w)
                except Exception: pass
            with open(out, "wb") as fh: w.write(fh)
            self.progress.emit(100)
            self.finished.emit(True, f"Bates-stamped {counter - start} pages", str(out))
        else:
            od = out if out.is_dir() else out.parent / (out.stem + "_bates")
            od.mkdir(parents=True, exist_ok=True)
            counter = start; n = len(files)
            for fi, f in enumerate(files):
                r = PdfReader(f); w = PdfWriter()
                w.append(r, import_outline=True)
                preserve_pdf_metadata_and_structure(r, w)
                for page in w.pages:
                    stamp(f"{prefix}{str(counter).zfill(digits)}{suffix}", page); counter += 1
                with open(od / Path(f).name, "wb") as fh: w.write(fh)
                self.progress.emit(int((fi+1)/n*100))
            self.finished.emit(True, f"Bates-stamped {counter - start} pages across {n} files", str(od))

    # ── Compare ──────────────────────────────────────────────
    def _run_compare(self):
        a = self.p["pdf_a"]; b = self.p["pdf_b"]
        out_h = Path(self.p["out_file"]); out_h.parent.mkdir(parents=True, exist_ok=True)
        iw = self.p.get("ignore_whitespace", True); ic = self.p.get("ignore_case", False)
        def norm(t):
            if iw: t = re.sub(r"\s+", " ", t).strip()
            if ic: t = t.lower()
            return t
        pa = [norm(p.extract_text() or "") for p in PdfReader(a).pages]
        self.progress.emit(25)
        pb = [norm(p.extract_text() or "") for p in PdfReader(b).pages]
        self.progress.emit(50)
        mp = max(len(pa), len(pb))
        reports = []; changed = 0
        for i in range(mp):
            ta = pa[i] if i < len(pa) else ""
            tb = pb[i] if i < len(pb) else ""
            if ta == tb:
                reports.append({"page": i+1, "status": "identical", "html": ""}); continue
            changed += 1
            diff = difflib.HtmlDiff(wrapcolumn=90)
            reports.append({"page": i+1, "status": "changed",
                            "html": diff.make_table(ta.splitlines(), tb.splitlines(),
                                    fromdesc=f"A — Page {i+1}", todesc=f"B — Page {i+1}",
                                    context=True, numlines=3)})
            self.progress.emit(50 + int((i+1)/mp*40))
        rows = ""
        for pr in reports:
            badge = "identical" if pr["status"] == "identical" else "changed"
            color = "#16a34a" if badge == "identical" else ACCENT
            rows += f'<tr><td>Page {pr["page"]}</td><td style="color:{color};font-weight:700">{badge}</td></tr>'
        details = ""
        for pr in reports:
            if pr["status"] == "identical": continue
            details += f'<h2 style="color:{ACCENT}">Page {pr["page"]}</h2>{pr["html"]}'
        html = f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><title>PDF Comparison</title>
<style>
body {{ font-family:'Segoe UI',Inter,sans-serif; background:#f8fafc; color:#0f172a; margin:0; }}
header {{ background:linear-gradient(135deg,{ACCENT} 0%,#60a5fa 100%); color:#fff; padding:32px 48px; }}
main {{ max-width:1200px; margin:32px auto; padding:0 24px 60px; }}
.summary {{ background:#fff; border:1px solid #e2e8f0; border-radius:12px; padding:20px; margin-bottom:24px; box-shadow:0 2px 12px rgba(0,0,0,.05); }}
.summary table {{ border-collapse:collapse; width:100%; }}
.summary td, .summary th {{ padding:8px 12px; border-bottom:1px solid #f1f5f9; text-align:left; }}
.summary th {{ background:#f8fafc; }}
.diff_table, .diff-table {{ font-family:"Cascadia Mono",Consolas,monospace; font-size:12px; border-collapse:collapse; margin:12px 0; width:100%; background:#fff; }}
td.diff_header {{ background:#f8fafc; color:#64748b; font-weight:700; padding:4px 8px; }}
.diff_add {{ background:#dcfce7; color:#166534; }}
.diff_chg {{ background:#fef3c7; color:#92400e; }}
.diff_sub {{ background:#fee2e2; color:#991b1b; }}
footer {{ text-align:center; color:#94a3b8; padding:20px 0; }}
</style></head><body>
<header><h1 style="font-size:1.8rem;font-weight:800;margin:0">PDF Comparison Report</h1>
<p style="margin-top:6px;opacity:.9"><b>A:</b> {Path(a).name} &nbsp;&nbsp; <b>B:</b> {Path(b).name}</p></header>
<main><div class="summary"><h2 style="margin-top:0">Summary</h2>
<p><b>{mp}</b> pages compared &mdash; <b style="color:{ACCENT}">{changed}</b> changed, <b style="color:#16a34a">{mp - changed}</b> identical.</p>
<table><tr><th>Page</th><th>Status</th></tr>{rows}</table></div>
{details if details else '<p style="color:#64748b;font-style:italic">No differences found.</p>'}
</main><footer>Generated by Document Studio Pro</footer></body></html>"""
        out_h.write_text(html, encoding="utf-8")
        self.progress.emit(100)
        self.finished.emit(True, f"Comparison report: {changed} page(s) differ", str(out_h))

    def _run_organize_visual(self):
        pdf_path = self.p["pdf_path"]; out = Path(self.p["out_file"])
        out.parent.mkdir(parents=True, exist_ok=True)
        order = self.p["page_order"]
        r = PdfReader(pdf_path); w = PdfWriter()
        kept = 0; total = len(order)
        for i, item in enumerate(order):
            if not item.get("keep", True): continue
            src = item["src"]
            if not (0 <= src < len(r.pages)): continue
            page = r.pages[src]
            rot = item.get("rot", 0) % 360
            if rot:
                try: page.rotate(rot)
                except Exception: pass
            w.add_page(page); kept += 1
            self.progress.emit(int((i+1)/total*95))
        preserve_pdf_metadata_and_structure(r, w)
        with open(out, "wb") as f: w.write(f)
        self.progress.emit(100)
        removed = total - kept
        msg = f"Rebuilt PDF: {kept} pages kept" + (f", {removed} removed" if removed else "")
        self.finished.emit(True, msg, str(out))

    # ── Phase 3 handlers ─────────────────────────────────────
    def _run_ocr(self):
        if fitz is None: raise RuntimeError("PyMuPDF required.")
        if pytesseract is None: raise RuntimeError("pytesseract not installed.")
        tess = find_tesseract()
        if not tess: raise RuntimeError("Tesseract binary not found.")
        pytesseract.pytesseract.tesseract_cmd = tess
        in_p = self.p["pdf_path"]; out = Path(self.p["out_file"])
        out.parent.mkdir(parents=True, exist_ok=True)
        lang = self.p.get("lang", "eng"); dpi = int(self.p.get("dpi", 200))
        self.log_line.emit(f"Tesseract: {tess}")
        src = fitz.open(in_p); total = len(src)
        out_doc = fitz.open()
        try:
            for i in range(total):
                page = src[i]
                pix = page.get_pixmap(dpi=dpi, alpha=False)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                try: pdf_bytes = pytesseract.image_to_pdf_or_hocr(img, extension="pdf", lang=lang)
                except Exception: pdf_bytes = pytesseract.image_to_pdf_or_hocr(img, extension="pdf")
                with fitz.open("pdf", pdf_bytes) as ocr:
                    out_doc.insert_pdf(ocr)
                self.log_line.emit(f"  Page {i+1}/{total} OCR'd")
                self.progress.emit(int((i+1)/total*95))
            out_doc.save(str(out), garbage=4, deflate=True, clean=True)
            out_doc.close(); src.close()
            self.progress.emit(100)
            self.finished.emit(True, f"OCR complete — {total} pages searchable ({lang})", str(out))
        except Exception as ex:
            raise RuntimeError(f"OCR failed: {ex}")

    def _run_verify_signatures(self):
        if not PYHANKO_OK: raise RuntimeError("pyhanko not installed.")
        pdf = self.p["pdf_path"]
        with open(pdf, "rb") as fh:
            reader = PyHankoReader(fh)
            sigs = reader.embedded_signatures
            if not sigs:
                self.finished.emit(True, "No signatures found.", pdf); return
            self.log_line.emit(f"\nFound {len(sigs)} signature(s):\n")
            for i, sig in enumerate(sigs, 1):
                self.log_line.emit(f"{'='*60}\n  SIGNATURE #{i}\n{'='*60}")
                try: self.log_line.emit(f"  Signer: {sig.signer_cert.subject.human_friendly}")
                except Exception: pass
                try: self.log_line.emit(f"  Issuer: {sig.signer_cert.issuer.human_friendly}")
                except Exception: pass
                try: self.log_line.emit(f"  Field:  {sig.field_name}")
                except Exception: pass
                try:
                    vc = PyHankoVC(allow_fetching=False)
                    st = validate_pdf_signature(sig, vc)
                    self.log_line.emit(f"  Integrity: {'INTACT' if st.intact else 'BROKEN'}")
                    self.log_line.emit(f"  Trusted:   {'YES' if st.trusted else 'No'}")
                    self.log_line.emit(f"  Valid:     {'YES' if st.valid else 'No'}")
                except Exception as ex: self.log_line.emit(f"  Validation: {ex}")
            self.progress.emit(100)
            self.finished.emit(True, f"{len(sigs)} signature(s) found", pdf)

    def _run_sign_pdf(self):
        if not PYHANKO_OK: raise RuntimeError("pyhanko not installed.")
        in_p = self.p["pdf_path"]; out = Path(self.p["out_file"])
        out.parent.mkdir(parents=True, exist_ok=True)
        pfx = self.p["pfx_path"]; pwd = self.p["pfx_password"]
        reason = self.p.get("reason", "Digitally signed")
        location = self.p.get("location", "")
        fname = self.p.get("field_name", "Signature1")
        if not os.path.exists(pfx): raise RuntimeError(f"Signing file not found: {pfx}")
        signer = pyhanko_signers.SimpleSigner.load_pkcs12(pfx_file=pfx, passphrase=pwd.encode())
        if signer is None: raise RuntimeError("Could not load PKCS#12 signer.")
        meta = PdfSignatureMetadata(field_name=fname, reason=reason, location=location)
        ps = PdfSigner(meta, signer=signer)
        with open(in_p, "rb") as fi, open(out, "wb") as fo:
            ps.sign_pdf(fi, output=fo)
        self.progress.emit(100)
        self.finished.emit(True, "PDF digitally signed", str(out))

    def _run_pdfa(self):
        gs = find_ghostscript()
        if not gs: raise RuntimeError("Ghostscript not found.")
        in_p = self.p["pdf_path"]; out = Path(self.p["out_file"])
        out.parent.mkdir(parents=True, exist_ok=True)
        lvl = self.p.get("level", "2")
        cmd = [gs, f"-dPDFA={lvl}", "-dBATCH", "-dNOPAUSE", "-dQUIET", "-dSAFER",
               "-sDEVICE=pdfwrite", "-dPDFACompatibilityPolicy=1",
               "-sColorConversionStrategy=UseDeviceIndependentColor",
               f"-sOutputFile={str(out)}", str(in_p)]
        try: res = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
        except subprocess.TimeoutExpired: raise RuntimeError("Ghostscript timed out.")
        if res.returncode != 0:
            raise RuntimeError(f"Ghostscript error: {(res.stderr or '').strip()[:300]}")
        if not out.exists() or out.stat().st_size == 0:
            raise RuntimeError("Ghostscript produced no output.")
        self.progress.emit(100)
        self.finished.emit(True, f"Converted to PDF/A-{lvl}b ({out.stat().st_size/1024:.0f} KB)", str(out))

    def _run_mail_merge(self):
        if DocxTemplate is None: raise RuntimeError("docxtpl not installed.")
        if Document is None: raise RuntimeError("python-docx not installed.")
        tpl = self.p["template"]; data = self.p["data_source"]
        od = Path(self.p["output_dir"]); od.mkdir(parents=True, exist_ok=True)
        fmt = self.p.get("format", "docx")
        nf = self.p.get("name_field", "").strip()
        if not os.path.exists(tpl): raise RuntimeError(f"Template not found: {tpl}")
        rows = []; ext = Path(data).suffix.lower()
        if ext == ".csv":
            with open(data, "r", encoding="utf-8-sig", newline="") as f:
                rows = list(csv.DictReader(f))
        elif ext == ".json":
            raw = json.loads(Path(data).read_text(encoding="utf-8"))
            rows = [raw] if isinstance(raw, dict) else raw
        elif ext in (".xlsx", ".xls"):
            if not (pythoncom and win32com):
                raise RuntimeError("Excel data requires pywin32. Export as .csv instead.")
            pythoncom.CoInitialize()
            xl = win32com.client.DispatchEx("Excel.Application")
            xl.Visible = False; xl.DisplayAlerts = False
            wb = xl.Workbooks.Open(os.path.abspath(data), ReadOnly=True)
            sheet = wb.Sheets(1); used = sheet.UsedRange
            nr = used.Rows.Count; nc = used.Columns.Count
            hdrs = [str(sheet.Cells(1, c).Value or f"col{c}") for c in range(1, nc+1)]
            for r in range(2, nr+1):
                row = {}
                for c in range(1, nc+1):
                    v = sheet.Cells(r, c).Value
                    row[hdrs[c-1]] = "" if v is None else str(v)
                rows.append(row)
            wb.Close(False); xl.Quit(); pythoncom.CoUninitialize()
        else:
            raise RuntimeError(f"Unsupported data source: {ext}")
        if not rows: raise RuntimeError("Data source has no rows.")
        self.log_line.emit(f"Loaded {len(rows)} row(s)")
        gen = 0; fail = 0
        for i, row in enumerate(rows):
            try:
                if nf and nf in row and str(row[nf]).strip():
                    base = re.sub(r'[^\w\-. ]', '_', str(row[nf]).strip())
                else:
                    base = f"document_{i+1:04d}"
                out_docx = od / f"{base}.docx"
                t = DocxTemplate(tpl); t.render(row); t.save(str(out_docx))
                if fmt == "pdf":
                    out_pdf = od / f"{base}.pdf"
                    if pythoncom and win32com:
                        try:
                            pythoncom.CoInitialize()
                            word = win32com.client.DispatchEx("Word.Application")
                            word.Visible = False; word.DisplayAlerts = 0
                            doc = word.Documents.Open(os.path.abspath(str(out_docx)), ReadOnly=True)
                            doc.ExportAsFixedFormat(OutputFileName=os.path.abspath(str(out_pdf)), ExportFormat=17, OpenAfterExport=False)
                            doc.Close(False); word.Quit(); pythoncom.CoUninitialize()
                            try: out_docx.unlink()
                            except Exception: pass
                        except Exception: pass
                gen += 1
                self.log_line.emit(f"  [{i+1}/{len(rows)}] -> {base}")
            except Exception as ex:
                fail += 1; self.log_line.emit(f"  FAILED row {i+1}: {ex}")
            self.progress.emit(int((i+1)/len(rows)*100))
        msg = f"Mail merge: {gen} generated" + (f", {fail} failed" if fail else "")
        self.finished.emit(gen > 0, msg, str(od))

    # ── Style / Source ───────────────────────────────────────
    def _run_style_inspect(self):
        doc_path = self.p["doc_path"]; out_report = self.p.get("out_report", "").strip()
        html_export = self.p.get("html_export", False)
        ext = Path(doc_path).suffix.lower(); fname = Path(doc_path).name
        def hr(t=""):
            sep = "─" * 68
            return f"\n{sep}\n  {t}\n{sep}" if t else f"\n{sep}"
        def pt(sz):
            try: return f"{sz.pt:.1f} pt" if sz else "(not set)"
            except Exception: return str(sz)
        lines = []; add = lines.append
        self.log_line.emit(f"Inspecting: {fname}"); self.progress.emit(10)
        if ext in (".docx", ".docm"):
            if not Document: raise RuntimeError("python-docx not installed.")
            doc = Document(doc_path)
            add(hr("DOCUMENT OVERVIEW"))
            add(f"  File name  : {fname}")
            add(f"  Paragraphs : {len(doc.paragraphs)}")
            add(f"  Tables     : {len(doc.tables)}")
            self.progress.emit(30)
            add(hr("STYLES"))
            for s in doc.styles:
                try:
                    fn = s.font
                    name = s.name
                    f_name = fn.name or "(inherited)"
                    f_size = pt(fn.size)
                    flags = []
                    if fn.bold: flags.append("B")
                    if fn.italic: flags.append("I")
                    if fn.underline: flags.append("U")
                    add(f"  • {name}  |  {f_name} @ {f_size}  [{'/'.join(flags) or '-'}]")
                except Exception: pass
            self.progress.emit(80)
        elif ext == ".pdf":
            r = PdfReader(doc_path); total = len(r.pages)
            add(hr("PDF OVERVIEW")); add(f"  Pages: {total}")
            fonts = set()
            for i, p in enumerate(r.pages):
                try:
                    res = p.get('/Resources')
                    if res and '/Font' in res:
                        for _, fo in res['/Font'].items():
                            try:
                                fd = fo.get_object() if hasattr(fo, 'get_object') else fo
                                fonts.add(str(fd.get('/BaseFont','?')).lstrip('/'))
                            except Exception: pass
                except Exception: pass
                self.progress.emit(int((i+1)/total*75) + 10)
            add(hr(f"FONTS — {len(fonts)}"))
            for f in sorted(fonts): add(f"  • {f}")
        else:
            add(f"Style inspection not supported for {ext}.")
        self.progress.emit(90)
        report = "\n".join(lines); self.log_line.emit(report)
        saved = ""
        if out_report:
            rp = Path(out_report); rp.parent.mkdir(parents=True, exist_ok=True)
            rp.write_text(f"Style Report\nDocument: {fname}\n" + "=" * 72 + "\n" + report, encoding="utf-8")
            saved = str(rp)
        self.progress.emit(100)
        if saved: self.finished.emit(True, "Style report saved", saved)
        else: self.finished.emit(True, "Style inspection complete — see log", doc_path)

    def _run_source_view(self):
        doc_path = self.p["doc_path"]; out_file = self.p.get("out_file", "").strip()
        component = self.p.get("component", "word/document.xml")
        page_num = self.p.get("page_num", 1); html_conv = self.p.get("html_convert", False)
        ext = Path(doc_path).suffix.lower()
        self.log_line.emit(f"Loading: {Path(doc_path).name}"); self.progress.emit(15)
        result = ""
        if html_conv and ext in (".docx", ".docm"):
            if not Document: raise RuntimeError("python-docx not installed.")
            doc = Document(doc_path); parts = []
            for bi, block in enumerate(doc.element.body):
                try:
                    tag = block.tag.split('}')[-1]
                    if tag == 'p':
                        from docx.text.paragraph import Paragraph as _P
                        p = _P(block, doc)
                        inner = "".join((r.text.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;') or '') for r in p.runs)
                        parts.append(f"<p>{inner or '&nbsp;'}</p>")
                    elif tag == 'tbl':
                        from docx.table import Table as _T
                        tbl = _T(block, doc); rows = ""
                        for ri, row in enumerate(tbl.rows):
                            cells = ""
                            for cell in row.cells:
                                tg = 'th' if ri == 0 else 'td'
                                ci = "".join(f"<p>{p.text}</p>" for p in cell.paragraphs)
                                cells += f"<{tg}>{ci}</{tg}>"
                            rows += f"<tr>{cells}</tr>"
                        parts.append(f"<table>{rows}</table>")
                except Exception: pass
                if bi % 20 == 0: self.progress.emit(35 + int(bi/max(len(doc.element.body),1)*45))
            result = f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><title>{Path(doc_path).stem}</title>
<style>body{{font-family:Calibri,sans-serif;font-size:11pt;line-height:1.6;color:#0f172a;background:#f8fafc;margin:0}}.doc-page{{max-width:820px;margin:40px auto;background:#fff;padding:72px 80px;box-shadow:0 0 40px rgba(0,0,0,.08);border-radius:4px}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #cbd5e1;padding:6px 10px}}th{{background:#f1f5f9}}</style>
</head><body><div class="doc-page">{''.join(parts)}</div></body></html>"""
            if not out_file: out_file = str(Path(doc_path).parent / f"{Path(doc_path).stem}_export.html")
        elif ext in (".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm"):
            with zipfile.ZipFile(doc_path, 'r') as z:
                ap = z.namelist()
                self.log_line.emit(f"Package contains {len(ap)} parts:")
                for p_ in ap: self.log_line.emit(f"    • {p_}")
                self.progress.emit(35)
                if component == "__list__": result = "\n".join(ap)
                else:
                    matches = [n for n in ap if component in n]
                    if not matches: matches = [n for n in ap if component.lower().split('/')[-1] in n.lower()]
                    if matches:
                        raw = z.read(matches[0])
                        try:
                            xml = raw.decode('utf-8', errors='replace')
                            try:
                                root = ET.fromstring(xml); ET.indent(root, space='  ')
                                xml = ET.tostring(root, encoding='unicode')
                            except Exception: pass
                            result = xml
                            self.log_line.emit(xml[:8000])
                        except Exception as e: result = f"Binary: {e}"
                    else: result = f"Part '{component}' not found."
        elif ext == ".pdf":
            r = PdfReader(doc_path); total = len(r.pages)
            if component == "__list__":
                for i, p in enumerate(r.pages):
                    try: self.log_line.emit(f"  Page {i+1}: {float(p.mediabox.width):.0f}x{float(p.mediabox.height):.0f}pt")
                    except Exception: pass
            else:
                idx = min(max(1, page_num), total) - 1
                p = r.pages[idx]
                try:
                    d = p.get_contents()
                    if d is None: result = "(no content)"
                    else:
                        b = d.get_data() if hasattr(d, 'get_data') else (d if isinstance(d, (bytes, bytearray)) else b"")
                        try: result = b.decode('latin-1')
                        except Exception: result = repr(b)
                    self.log_line.emit(result[:8000])
                except Exception as e: result = f"Error: {e}"
        elif ext in (".html", ".htm", ".txt", ".md", ".rst"):
            with open(doc_path, 'r', encoding='utf-8', errors='replace') as f: result = f.read()
            self.log_line.emit(result[:8000])
        else:
            result = f"Not supported: {ext}"
        self.progress.emit(92)
        if out_file and result:
            op = Path(out_file); op.parent.mkdir(parents=True, exist_ok=True)
            op.write_text(result, encoding='utf-8')
            self.progress.emit(100); self.finished.emit(True, "Saved to file", str(op))
            if str(op).lower().endswith('.html'):
                import webbrowser
                try: webbrowser.open(op.as_uri())
                except Exception: pass
        else:
            self.progress.emit(100); self.finished.emit(True, "Source shown in log", doc_path)

